#!/usr/bin/env python3
"""Live Link Tracker: paste vendor live links -> download Excel.

Run:  python3 live_tracker.py      then open http://localhost:8765
Columns: SNO | PROFILE NAME | PROFILE LINK | FOLLOWERS | LIVE LINK | LIVE DATE
"""
import os, io, re, json, base64, urllib.request, urllib.parse
from datetime import datetime, timezone
from concurrent.futures import ThreadPoolExecutor
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

ALPHABET = "ABCDEFGHIJKLMNOPQRSTUVWXYZabcdefghijklmnopqrstuvwxyz0123456789-_"
UA = {"User-Agent": "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/605.1.15 "
                    "(KHTML, like Gecko) Version/17.0 Safari/605.1.15",
      "Accept-Language": "en-US,en;q=0.9"}
LINK_RE = re.compile(r"instagram\.com/(?:([A-Za-z0-9._]+)/)?(reel|reels|p|tv)/([A-Za-z0-9_-]+)")


def parse_link(raw):
    """Return (username|None, shortcode, clean_link) or None."""
    m = LINK_RE.search(raw.strip())
    if not m:
        return None
    user, kind, code = m.groups()
    if user in ("reel", "reels", "p", "tv"):  # no username in url
        user = None
    kind = "reel" if kind == "reels" else kind
    return user, code, f"https://www.instagram.com/{kind}/{code}/"


def date_from_shortcode(code):
    """Instagram media id embeds the post time (ms since 2011-08-24)."""
    try:
        n = 0
        for ch in code[:11]:
            n = n * 64 + ALPHABET.index(ch)
        ts = (n >> 23) + 1314220021721
        return datetime.fromtimestamp(ts / 1000, tz=timezone.utc).astimezone().strftime("%d-%m-%Y")
    except ValueError:
        return ""


def fetch(url):
    try:
        req = urllib.request.Request(url, headers=UA)
        return urllib.request.urlopen(req, timeout=15).read().decode("utf-8", "ignore")
    except Exception:
        return ""


def username_from_page(clean_link):
    html = fetch(clean_link)
    for pat in (r'"owner":\{[^}]*?"username":"([A-Za-z0-9._]+)"',
                r'instagram\.com/([A-Za-z0-9._]+)/(?:reel|p)/',
                r'\(@([A-Za-z0-9._]+)\)'):
        m = re.search(pat, html)
        if m and m.group(1) not in ("reel", "p"):
            return m.group(1)
    return None


def to_number(s):
    s = s.replace(",", "").strip()
    mult = {"K": 1e3, "M": 1e6, "B": 1e9}.get(s[-1:].upper(), 1)
    return int(float(s[:-1] if mult != 1 else s) * mult)


def followers_of(username):
    html = fetch(f"https://www.instagram.com/{username}/")
    m = re.search(r'content="([\d.,]+[KMB]?) Followers', html, re.I)
    return to_number(m.group(1)) if m else None


def build_rows(text, followers=True):
    def one(p):
        user, code, clean = p
        user = user or username_from_page(clean)
        return {
            "name": user or "",
            "plink": f"https://www.instagram.com/{user}/" if user else "",
            "followers": followers_of(user) if (user and followers) else None,
            "live": clean,
            "date": date_from_shortcode(code),
        }
    parsed = [parse_link(l) for l in re.split(r"[\s,]+", text) if l]
    with ThreadPoolExecutor(6) as ex:
        rows = list(ex.map(one, [p for p in parsed if p]))
    counts = {}
    for r in rows:
        counts[r["live"]] = counts.get(r["live"], 0) + 1
    for r in rows:
        r["dup"] = counts[r["live"]] > 1
    return rows


def make_xlsx(rows):
    wb = Workbook(); ws = wb.active; ws.title = "Live Links"
    ws.append(["SNO", "PROFILE NAME", "PROFILE LINK", "FOLLOWERS", "LIVE LINK", "LIVE DATE"])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F3864")
        c.alignment = Alignment(horizontal="center")
    gap = PatternFill("solid", fgColor="FFF2CC")  # yellow = fill manually
    for i, r in enumerate(rows, 1):
        ws.append([i, r["name"], r["plink"], r["followers"], r["live"], r["date"]])
        for col in (2, 3, 4):
            if ws.cell(i + 1, col).value in (None, ""):
                ws.cell(i + 1, col).fill = gap
        if r.get("dup"):  # duplicate live link -> whole row red
            for col in range(1, 7):
                c = ws.cell(i + 1, col)
                c.fill = PatternFill("solid", fgColor="FF9999")
                c.font = Font(color="9C0006", bold=True)
        if r["followers"]:
            ws.cell(i + 1, 4).number_format = "#,##0"
    for col, w in zip("ABCDEF", (6, 26, 40, 14, 48, 14)):
        ws.column_dimensions[col].width = w
    buf = io.BytesIO(); wb.save(buf); return buf.getvalue()


def fill_plan(xlsx_bytes, text):
    """Paste live link + date next to each page in an uploaded plan."""
    wb = load_workbook(io.BytesIO(xlsx_bytes)); ws = wb.active
    heads = {str(c.value).strip().upper(): c.column for c in ws[1] if c.value}
    name_col = next((v for k, v in heads.items() if "NAME" in k), None)
    link_col = next((v for k, v in heads.items() if "LINK" in k), None)
    by_user = {}
    for r in range(2, ws.max_row + 1):
        keys = []
        if name_col and ws.cell(r, name_col).value:
            keys.append(str(ws.cell(r, name_col).value))
        if link_col and ws.cell(r, link_col).value:
            m = re.search(r"instagram\.com/([A-Za-z0-9._]+)", str(ws.cell(r, link_col).value))
            if m: keys.append(m.group(1))
        for k in keys:
            by_user.setdefault(k.strip().lower().lstrip("@"), r)
    lc, dc = ws.max_column + 1, ws.max_column + 2
    hdr = ws.cell(1, 1)
    for col, t in ((lc, "LIVE LINK"), (dc, "LIVE DATE")):
        c = ws.cell(1, col, t); c.font = hdr.font.copy(); c.fill = hdr.fill.copy(); c.alignment = hdr.alignment.copy()
    ws.column_dimensions[ws.cell(1, lc).column_letter].width = 48
    ws.column_dimensions[ws.cell(1, dc).column_letter].width = 14
    red = PatternFill("solid", fgColor="FF9999"); redf = Font(color="9C0006", bold=True)
    rows = build_rows(text, followers=False)
    seen, extra = {}, []
    for x in rows:
        r = by_user.get(x["name"].lower())
        if x["live"] in seen:  # duplicate live link -> red
            for rr in {seen[x["live"]], r} - {None}:
                for col in (lc, dc):
                    ws.cell(rr, col).fill = red; ws.cell(rr, col).font = redf
            extra.append((x["name"], x["live"], x["date"], "DUPLICATE LINK"))
            continue
        if r is None:
            extra.append((x["name"], x["live"], x["date"], "Page not found in plan" if x["name"] else "Could not read username"))
        elif ws.cell(r, lc).value:
            ws.cell(r, lc).fill = red; ws.cell(r, lc).font = redf
            extra.append((x["name"], x["live"], x["date"], "Page already has a live link (2nd link)"))
        else:
            ws.cell(r, lc, x["live"]); ws.cell(r, dc, x["date"])
            seen[x["live"]] = r
    if extra:
        e = wb.create_sheet("Not pasted")
        e.append(["PAGE NAME", "LIVE LINK", "LIVE DATE", "REASON"])
        for row in extra: e.append(row)
        for col, w in zip("ABCD", (26, 48, 14, 40)): e.column_dimensions[col].width = w
        for c in e[1]: c.font = Font(bold=True)
    buf = io.BytesIO(); wb.save(buf)
    return buf.getvalue(), len(rows) - len(extra), len(extra)


PAGE = b"""<!doctype html><meta charset=utf-8><title>Live Link Tracker</title>
<meta name=viewport content="width=device-width,initial-scale=1">
<body style="font:16px system-ui;max-width:760px;margin:40px auto;padding:0 16px">
<h2>Live Link Tracker</h2>
<p>Paste the vendor's live links (one per line). Anything after <code>?</code> is removed. Duplicates are marked red.</p>
<textarea id=links rows=12 style="width:100%;font:14px monospace" placeholder="https://www.instagram.com/reel/XXXX/?igsh=..."></textarea>
<p><b>Option A &ndash; New tracker</b><br>
<button onclick="go('tracker')" style="padding:10px 20px;font-size:16px">Download tracker Excel</button></p>
<p><b>Option B &ndash; Paste into an existing plan</b><br>
<input type=file id=plan accept=".xlsx"> 
<button onclick="go('plan')" style="padding:10px 20px;font-size:16px">Fill plan with live links</button></p>
<p id=msg style="color:#444"></p>
<script>
async function go(mode){
  const msg=document.getElementById('msg'); const links=document.getElementById('links').value;
  if(!links.trim()){msg.textContent='Paste some links first.';return}
  let file='';
  if(mode=='plan'){
    const f=document.getElementById('plan').files[0];
    if(!f){msg.textContent='Choose the plan .xlsx first.';return}
    file=await new Promise(r=>{const fr=new FileReader();fr.onload=()=>r(fr.result.split(',')[1]);fr.readAsDataURL(f)});
  }
  msg.textContent='Working... (looking up pages, can take a minute)';
  const res=await fetch('/'+mode,{method:'POST',body:JSON.stringify({links,file})});
  if(!res.ok){msg.textContent='Error: '+await res.text();return}
  const blob=await res.blob(); const a=document.createElement('a');
  a.href=URL.createObjectURL(blob); a.download=mode=='plan'?'plan_with_live_links.xlsx':'live_links.xlsx'; a.click();
  msg.textContent='Done. '+(res.headers.get('X-Summary')||'');
}
</script></body>"""


class H(BaseHTTPRequestHandler):
    def do_GET(self):
        self.send_response(200); self.send_header("Content-Type", "text/html"); self.end_headers()
        self.wfile.write(PAGE)

    def do_POST(self):
        try:
            n = int(self.headers["Content-Length"])
            if n > 8_000_000: raise ValueError("File too large (max ~5 MB)")
            d = json.loads(self.rfile.read(n))
            if len(re.findall(r"instagram\.com", d["links"])) > 300: raise ValueError("Max 300 links per run")
            summary = ""
            if self.path == "/plan":
                data, ok, bad = fill_plan(base64.b64decode(d["file"]), d["links"])
                summary = f"{ok} pasted into plan, {bad} listed on the 'Not pasted' sheet."
            else:
                data = make_xlsx(build_rows(d["links"]))
        except Exception as e:
            self.send_response(500); self.end_headers(); self.wfile.write(str(e).encode()); return
        self.send_response(200)
        self.send_header("Content-Type", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        self.send_header("X-Summary", summary)
        self.end_headers(); self.wfile.write(data)


if __name__ == "__main__":
    port = int(os.environ.get("PORT", 8765))
    host = "0.0.0.0" if os.environ.get("PORT") else "127.0.0.1"
    print(f"Open http://localhost:{port}")
    ThreadingHTTPServer((host, port), H).serve_forever()
